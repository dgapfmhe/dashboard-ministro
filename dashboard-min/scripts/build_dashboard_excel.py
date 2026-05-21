"""Construye un Excel profesional con gráficos nativos manipulables.

Replica las pestañas del dashboard como hojas de Excel con sus gráficos
referenciados a los datos. Al abrir en Excel:
  - Las hojas `data_*` contienen los datos crudos en formato ancho.
  - Las hojas `tab_*` tienen gráficos nativos (Bar/Line/Combo) que apuntan
    a rangos de las hojas `data_*`. Modificar un valor en `data_*` actualiza
    el gráfico automáticamente.
  - El usuario puede cambiar colores, escalas, agregar series, etc.

Salida: `../dashboard-web/public/dashboard_completo.xlsx`

Ejecutar:
    python scripts/build_dashboard_excel.py
"""
from __future__ import annotations

import os
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings
from supabase import create_client

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
OUTPUT = Path(__file__).resolve().parent.parent.parent / "dashboard-web" / "public" / "dashboard_completo.xlsx"


# ---------------------------------------------------------------------------
# Fetch series de Supabase (paginado)
# ---------------------------------------------------------------------------
def fetch_all_series() -> dict[str, list[tuple[str, float]]]:
    """Devuelve {variable_id: [(fecha, valor), ...]}, ordenado por fecha."""
    s = get_settings()
    client = create_client(s.supabase_url, s.supabase_service_key)
    PAGE = 1000
    by_var: dict[str, list[tuple[str, float]]] = defaultdict(list)
    offset = 0
    while True:
        res = (client.table("series_datos")
               .select("variable_id,fecha,valor")
               .order("variable_id")
               .order("fecha")
               .range(offset, offset + PAGE - 1)
               .execute())
        rows = res.data or []
        if not rows:
            break
        for r in rows:
            if r["valor"] is None:
                continue
            try:
                v = float(r["valor"])
                if v != v or v in (float("inf"), float("-inf")):
                    continue
            except (TypeError, ValueError):
                continue
            by_var[r["variable_id"]].append((r["fecha"], v))
        if len(rows) < PAGE:
            break
        offset += PAGE
    print(f"Series Supabase: {len(by_var)} variables, {sum(len(v) for v in by_var.values())} filas totales")
    return dict(by_var)


def fetch_catalog() -> list[dict]:
    s = get_settings()
    client = create_client(s.supabase_url, s.supabase_service_key)
    res = (client.table("variables")
           .select("id,nombre,modulo,submodulo,frecuencia,unidad,escala,fuente,tipo,formula")
           .order("modulo").order("id").execute())
    return res.data or []


# ---------------------------------------------------------------------------
# Normalización de fechas
# ---------------------------------------------------------------------------
def normalize_fecha(fecha: str, freq: str) -> str:
    if freq == "diaria":
        return fecha[:10]
    if freq == "mensual":
        return fecha[:7] + "-01"
    if freq == "trimestral":
        y, m, *_ = fecha.split("-")
        mn = int(m); qm = ((mn - 1) // 3) * 3 + 1
        return f"{y}-{qm:02d}-01"
    if freq == "anual":
        return fecha[:4] + "-01-01"
    return fecha


# ---------------------------------------------------------------------------
# Estilos MÍNIMOS (Excel se quejaba de "Repaired Records" con PatternFill +
# Font custom; ahora solo bold sin fill ni colores hex propios)
# ---------------------------------------------------------------------------
HDR_FONT = Font(bold=True)


def fmt_header(ws, row: int, n_cols: int):
    """Solo aplica bold a la primera fila — sin fills ni colores custom para
    evitar que Excel rechace los estilos en styles.xml."""
    for c in range(1, n_cols + 1):
        ws.cell(row=row, column=c).font = HDR_FONT


def autosize_cols(ws, widths: dict[int, int]):
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


# ---------------------------------------------------------------------------
# Construir hoja de datos ANCHA (formato pivot) por frecuencia
# ---------------------------------------------------------------------------
def build_data_sheet(wb: Workbook, sheet_name: str,
                     vars_in_sheet: list[dict], series: dict[str, list[tuple[str, float]]],
                     freq: str) -> dict[str, int]:
    """Crea hoja con datos en formato ancho. Retorna `{var_id: col_idx}`.

    LAYOUT (clave para que los charts no rompan):
      - Fila 1: cabecera (variable_id), columna A = "fecha"
      - Fila 2+: datos (sin subheader de texto entre cabecera y datos)
    De esta forma `Reference(..., min_row=1, max_row=N)` con
    `titles_from_data=True` interpreta correctamente fila 1 como título.
    """
    date_set: set[str] = set()
    var_data: dict[str, dict[str, float]] = {}
    for v in vars_in_sheet:
        vid = v["id"]
        m: dict[str, float] = {}
        for f, val in series.get(vid, []):
            fn = normalize_fecha(f, freq)
            m[fn] = val
            date_set.add(fn)
        var_data[vid] = m
    dates = sorted(date_set)
    if not dates:
        return {}

    ws = wb.create_sheet(sheet_name)
    # ─── Fila 1: cabecera única ───
    ws.cell(row=1, column=1, value="fecha")
    col_idx: dict[str, int] = {}
    for j, v in enumerate(vars_in_sheet, start=2):
        ws.cell(row=1, column=j, value=v["id"])
        col_idx[v["id"]] = j
    fmt_header(ws, 1, 1 + len(vars_in_sheet))

    # ─── Filas 2+: datos puros ───
    for ri, fecha in enumerate(dates, start=2):
        ws.cell(row=ri, column=1, value=fecha)
        for v in vars_in_sheet:
            val = var_data[v["id"]].get(fecha)
            if val is not None:
                ws.cell(row=ri, column=col_idx[v["id"]], value=val)

    ws.freeze_panes = "B2"
    widths = {1: 12}
    for v in vars_in_sheet:
        widths[col_idx[v["id"]]] = max(len(v["id"]) + 2, 12)
    autosize_cols(ws, widths)

    return col_idx


# ---------------------------------------------------------------------------
# Helpers de charts
# ---------------------------------------------------------------------------
def _fix_axes(chart):
    """openpyxl default: axPos='l' en ambos ejes → Excel marca archivo
    corrupto. Setear directamente para que el catAx vaya abajo y valAx
    a la izquierda. Setear `position` NO funciona, hay que usar `axPos`."""
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    chart.x_axis.delete = False
    chart.y_axis.delete = False


# Tamaños chicos para que se vean limpios en Excel y no se pisen
CHART_W = 14
CHART_H = 8

# Cuántos puntos mostrar en los charts (últimos N) según frecuencia.
WINDOW_BY_FREQ: dict[str, int] = {
    "diaria":     365 * 10,
    "mensual":    12 * 10,
    "trimestral": 4 * 10,
    "anual":      10,
}

# Colores fijos por variable_id — mismo color en TODOS los charts donde aparece.
# Asegura consistencia visual: 'tpm' siempre azul navy, 'embi_rd' siempre rojo, etc.
VAR_COLORS: dict[str, str] = {
    # Tasas
    "tpm":              "1B3A5C",  # navy
    "tia_prom_pond":    "D32F2F",  # rojo
    "tip_prom_pond":    "2E7D32",  # verde
    "tib_rd":           "F57C00",  # naranja
    "ti_fond_fed":      "7F7F7F",  # gris
    # Inflación
    "ipc_yoy":          "1B3A5C",
    "ipc_sub":          "D32F2F",
    "ipc_rd":           "1B3A5C",
    "cpi_usa":          "7F7F7F",
    # EMBI
    "embi_global":      "7F7F7F",  # gris
    "embi_latam":       "F57C00",  # ámbar
    "embi_rd":          "D32F2F",  # rojo
    "embi":             "9E9E9E",
    # Tipo de cambio
    "tcn_compra":       "2E7D32",
    "tcn_venta":        "1B3A5C",
    "tcn_venta_m":      "1B3A5C",
    "tcn_compra_m":     "2E7D32",
    # RIN / Externo
    "rin":              "1B3A5C",
    "exp_nom":          "2E7D32",
    "imp_nom":          "D32F2F",
    "rem_fam":          "F57C00",
    "ied_bp":           "9C27B0",  # morado
    "viajes_bp":        "2C7A7B",  # teal
    "se_cc_numtur":     "2C7A7B",
    # Distribuidoras eléctricas
    "el_gen_total":     "1B3A5C",
    "el_gen_carbon":    "F57C00",
    "el_gen_gas_nat":   "1976D2",  # azul claro
    "el_gen_fo2":       "D32F2F",
    "el_gen_fo6":       "A0522D",  # marrón
    "el_gen_hidro":     "2E7D32",
    "el_gen_eolica":    "00838F",
    "el_gen_solar":     "F9A825",  # amarillo
    "el_gen_biomasa":   "388E3C",
    "el_perd_tot":      "1B3A5C",
    "el_perd_edn":      "2E7D32",
    "el_perd_eds":      "F57C00",
    "el_perd_ede":      "D32F2F",
    "el_cobr_tot":      "1B3A5C",
    "el_cobr_edn":      "2E7D32",
    "el_cobr_eds":      "F57C00",
    "el_cobr_ede":      "D32F2F",
    "el_pcomb_fo2":     "D32F2F",
    "el_pcomb_fo6":     "A0522D",
    "el_pcomb_gas":     "1976D2",
    "el_pcomb_carbon":  "F57C00",
    # Préstamos por sector
    "prest_agr":        "2E7D32",
    "prest_ind":        "D32F2F",
    "prest_con":        "F57C00",
    "prest_com":        "1976D2",
    "prest_viv":        "9C27B0",
    "prest_cons":       "9E9E9E",
    "prest_tar":        "F9A825",
    # OSD
    "aen_osd_act":      "2E7D32",
    "aen_osd_pas":      "D32F2F",
    "dep_sp_osd":       "1B3A5C",
    "prest_me_usd_osd": "1B3A5C",
    # Otros
    "bmr_rd":           "1B3A5C",
    "deprem_cp":        "1B3A5C",
}


def _apply_color(series, vid: str):
    """Setea color de línea/borde según VAR_COLORS para consistencia visual."""
    color = VAR_COLORS.get(vid)
    if not color:
        return
    try:
        from openpyxl.chart.shapes import GraphicalProperties
        from openpyxl.drawing.line import LineProperties
        from openpyxl.drawing.fill import ColorChoice
        series.graphicalProperties = GraphicalProperties(
            solidFill=color,
            ln=LineProperties(solidFill=color),
        )
    except Exception:
        pass


def _data_first_row(source_ws, col_idx: int, header_row: int = 1,
                    last_row: int | None = None) -> int | None:
    """Primera fila con valor en la columna `col_idx`. Devuelve None si vacía."""
    last_row = last_row or source_ws.max_row
    for r in range(header_row + 1, last_row + 1):
        if source_ws.cell(row=r, column=col_idx).value is not None:
            return r
    return None


def _window_rows_aligned(source_ws, freq: str, n_data_rows: int,
                          col_indices: list[int]) -> tuple[int, int]:
    """Devuelve (start_row, end_row) del chart de manera que TODAS las
    variables referenciadas tengan datos desde el inicio (sin NA al principio).

    Lógica:
      - end_row = última fila de datos
      - window_start = end - WINDOW_BY_FREQ[freq] + 1  (ej: 10 años atrás)
      - latest_first = MAX(primera fila con dato) entre las variables
      - start_row = MAX(window_start, latest_first)
    """
    end_row = 1 + n_data_rows
    window = WINDOW_BY_FREQ.get(freq, n_data_rows)
    window_start = max(2, end_row - window + 1)
    latest_first = 2
    for c in col_indices:
        fr = _data_first_row(source_ws, c, header_row=1, last_row=end_row)
        if fr is None:
            continue
        latest_first = max(latest_first, fr)
    start_row = max(window_start, latest_first)
    return start_row, end_row


from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef


def _set_series_title(series, source_ws, col_letter: str):
    series.tx = SeriesLabel(strRef=StrRef(f"'{source_ws.title}'!${col_letter}$1"))


def add_line_chart(target_ws, source_ws, *, title: str, var_ids: list[str],
                   col_idx: dict[str, int], n_data_rows: int, anchor: str,
                   freq: str = "mensual"):
    if n_data_rows < 2:
        return
    # Filtrar variables que SÍ existen en este sheet
    valid_vars = [v for v in var_ids if v in col_idx]
    if not valid_vars:
        return
    col_indices = [col_idx[v] for v in valid_vars]
    # Ventana alineada: arranca donde TODAS las vars tienen datos
    start_data, end_data = _window_rows_aligned(source_ws, freq, n_data_rows, col_indices)
    if end_data - start_data < 2:
        return

    chart = LineChart()
    chart.title = title
    chart.width = CHART_W
    chart.height = CHART_H

    for vid in valid_vars:
        c = col_idx[vid]
        data_ref = Reference(source_ws, min_col=c, max_col=c,
                              min_row=start_data, max_row=end_data)
        chart.add_data(data_ref, titles_from_data=False)
        _set_series_title(chart.series[-1], source_ws, get_column_letter(c))
        _apply_color(chart.series[-1], vid)

    cats_ref = Reference(source_ws, min_col=1, max_col=1,
                          min_row=start_data, max_row=end_data)
    chart.set_categories(cats_ref)
    _fix_axes(chart)
    target_ws.add_chart(chart, anchor)


def add_bar_chart(target_ws, source_ws, *, title: str, var_id: str,
                  col_idx: dict[str, int], n_data_rows: int, anchor: str,
                  freq: str = "mensual"):
    if var_id not in col_idx or n_data_rows < 2:
        return
    c = col_idx[var_id]
    start_data, end_data = _window_rows_aligned(source_ws, freq, n_data_rows, [c])
    if end_data - start_data < 2:
        return

    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.width = CHART_W
    chart.height = CHART_H
    chart.legend = None

    data_ref = Reference(source_ws, min_col=c, max_col=c,
                          min_row=start_data, max_row=end_data)
    cats_ref = Reference(source_ws, min_col=1, max_col=1,
                          min_row=start_data, max_row=end_data)
    chart.add_data(data_ref, titles_from_data=False)
    _set_series_title(chart.series[-1], source_ws, get_column_letter(c))
    _apply_color(chart.series[-1], var_id)
    chart.set_categories(cats_ref)
    _fix_axes(chart)
    target_ws.add_chart(chart, anchor)


def add_combo_chart(target_ws, source_ws, *, title: str, bar_var: str,
                    line_var: str, col_idx: dict[str, int], n_data_rows: int,
                    anchor: str, freq: str = "mensual"):
    if bar_var not in col_idx or line_var not in col_idx or n_data_rows < 2:
        return
    cb = col_idx[bar_var]
    cl = col_idx[line_var]
    start_data, end_data = _window_rows_aligned(source_ws, freq, n_data_rows, [cb, cl])
    if end_data - start_data < 2:
        return

    bc = BarChart()
    bc.type = "col"
    bc.title = title
    bc.width = CHART_W
    bc.height = CHART_H

    data_ref = Reference(source_ws, min_col=cb, max_col=cb,
                          min_row=start_data, max_row=end_data)
    cats_ref = Reference(source_ws, min_col=1, max_col=1,
                          min_row=start_data, max_row=end_data)
    bc.add_data(data_ref, titles_from_data=False)
    _set_series_title(bc.series[-1], source_ws, get_column_letter(cb))
    _apply_color(bc.series[-1], bar_var)
    bc.set_categories(cats_ref)
    _fix_axes(bc)

    lc = LineChart()
    line_ref = Reference(source_ws, min_col=cl, max_col=cl,
                          min_row=start_data, max_row=end_data)
    lc.add_data(line_ref, titles_from_data=False)
    _set_series_title(lc.series[-1], source_ws, get_column_letter(cl))
    _apply_color(lc.series[-1], line_var)
    lc.y_axis.axId = 200
    lc.y_axis.crosses = "max"
    lc.y_axis.axPos = "r"
    bc += lc
    target_ws.add_chart(bc, anchor)


# ---------------------------------------------------------------------------
# Hoja README
# ---------------------------------------------------------------------------
def build_readme(wb: Workbook, fecha_corte: str, n_vars: int, n_filas: int):
    ws = wb.create_sheet("README", 0)
    lines: list[tuple[str, bool]] = [
        ("Panel Macroeconómico DGAPF — Excel descargable", True),
        (f"Generado: {fecha_corte}", False),
        ("", False),
        ("CÓMO USAR ESTE ARCHIVO", True),
        ("Las hojas tab_* contienen gráficos nativos de Excel.", False),
        ("Cada gráfico está referenciado a la hoja data_* correspondiente.", False),
        ("Modificar un valor en data_* actualiza el gráfico automáticamente.", False),
        ("", False),
        ("HOJAS DE DATOS (referencia para los gráficos)", True),
        ("data_diaria       — series diarias (TCN, EMBI, commodities, depósitos remunerados)", False),
        ("data_mensual      — series mensuales (IPC, IMAE, RIN, agregados monetarios)", False),
        ("data_trimestral   — series trimestrales (PIB y componentes, BoP, deuda)", False),
        ("data_anual        — series anuales (balances financieros eléctricos)", False),
        ("Catalogo          — metadata de las 500+ variables", False),
        ("", False),
        ("HOJAS DE GRÁFICOS (réplica de las pestañas del dashboard)", True),
        ("tab_externo            — Sector Externo", False),
        ("tab_tasas              — Tasas", False),
        ("tab_tipo_cambio        — Tipo de Cambio", False),
        ("tab_depositos_prest    — Depósitos y Préstamos", False),
        ("tab_dep_remunerados    — Depósitos Remunerados CP", False),
        ("tab_sector_electrico   — Sector Eléctrico", False),
        ("", False),
        ("FÓRMULAS DE VARIABLES DERIVADAS (ver hoja 'formulas')", True),
        ("  TCR  = (TCN_venta_m/1.358727436) x (CPI_USA/218.076x100) / IPC_RD", False),
        ("  RIN%PIB  = RIN / (PIB_4Q_trailing / TCN_promedio) x 100", False),
        ("  Prestamos_ME%  = (prest_me_usd x tc_fin_mes) / prest_tot x 100", False),
        ("  Taylor  = r* + pi* + alpha(pi - pi*) + beta(y - y_pot)", False),
        ("  AEN_OSD = aen_osd_act - aen_osd_pas", False),
        ("", False),
        ("FUENTES", True),
        ("BCRD, FRED, Yahoo Finance, Banco Mundial CMO, MEM, Hacienda.", False),
        ("Pipeline corre diario (Yahoo/FRED) y semanal (BCRD).", False),
        ("", False),
        ("ESTADISTICAS DE ESTE ARCHIVO", True),
        (f"Variables con data: {n_vars}", False),
        (f"Filas totales: {n_filas:,}", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if bold:
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 100


def build_catalog(wb: Workbook, cat: list[dict]):
    ws = wb.create_sheet("Catalogo")
    headers = ["id", "nombre", "modulo", "submodulo", "frecuencia", "unidad", "escala", "fuente", "tipo", "formula"]
    ws.append(headers)
    fmt_header(ws, 1, len(headers))
    for v in cat:
        ws.append([v.get(h) for h in headers])
    ws.freeze_panes = "A2"
    autosize_cols(ws, {1: 22, 2: 50, 3: 18, 4: 14, 5: 14, 6: 14, 7: 12, 8: 14, 9: 10, 10: 50})


def build_formulas(wb: Workbook):
    ws = wb.create_sheet("formulas")
    headers = ["variable / cálculo", "descripción", "fórmula", "fuentes"]
    rows = [
        ("tcr_rd", "Tipo de Cambio Real bilateral RD-USA",
         "(tcn_venta_m / 1.358727436) × (cpi_usa / 218.076 × 100) / ipc_rd",
         "tcn_venta_m: BCRD · cpi_usa: FRED CPIAUCSL · ipc_rd: BCRD"),
        ("rin_pct_pib", "RIN sobre PIB anualizado",
         "rin / (Σ pib_nom[Q-3..Q0] / tcn_venta_m) × 100",
         "rin: BCRD · pib_nom: BCRD pib_gasto · tcn_venta_m: BCRD"),
        ("aen_osd", "Activos Externos Netos Sistema Financiero",
         "aen_osd_act − aen_osd_pas",
         "BCRD credito_osd.xlsx"),
        ("prest_me_pct", "Préstamos en ME como % del total",
         "(prest_me_usd_osd × tc_fin_mes_osd) / prest_tot_osd × 100",
         "BCRD credito_osd.xlsx"),
        ("taylor_rate", "Regla de Taylor",
         "i = r* + π* + α·(π − π*) + β·(y − ȳ)",
         "tpm/ipc_yoy/imae: BCRD"),
        ("yoy_growth", "Crecimiento interanual",
         "(x[t] / x[t-12] − 1) × 100",
         "cálculo en frontend / Excel"),
        ("ma12", "Media móvil 12 meses",
         "promedio(x[t-11..t])",
         "cálculo en frontend / Excel"),
        ("deprem_real", "Depósitos remunerados deflactados",
         "deprem_cp / (ipc_rd[mes] / ipc_rd[mes_base])",
         "BCRD operaciones_monetarias.xlsx + ipc_rd"),
        ("ied_bp (signo invertido)", "Inversión Extranjera Directa (entrada)",
         "× -1 sobre Cuenta Financiera del MBP6 → POSITIVO = entrada de IED",
         "BCRD bpagos_trim_6.xls fila 61"),
    ]
    ws.append(headers); fmt_header(ws, 1, len(headers))
    for r in rows: ws.append(r)
    ws.freeze_panes = "A2"
    autosize_cols(ws, {1: 28, 2: 50, 3: 70, 4: 50})


# ---------------------------------------------------------------------------
# Specs de tabs: qué chart va en qué hoja
# ---------------------------------------------------------------------------
TAB_SPECS = [
    {
        "sheet": "tab_externo",
        "title": "SECTOR EXTERNO",
        "charts": [
            {"kind": "bar",   "title": "RIN — Reservas Internacionales Netas (MM US$)",
             "freq": "mensual", "var": "rin", "anchor": "A2"},
            {"kind": "line",  "title": "EMBI — Global, LATAM, RD",
             "freq": "diaria", "vars": ["embi_global", "embi_latam", "embi_rd"], "anchor": "L2"},
            {"kind": "bar",   "title": "IED — Inversión Extranjera Directa (MM US$, BoP trim.)",
             "freq": "trimestral", "var": "ied_bp", "anchor": "A22"},
            {"kind": "bar",   "title": "Ingresos por Turismo (MM US$, Viajes BoP)",
             "freq": "trimestral", "var": "viajes_bp", "anchor": "L22"},
            {"kind": "bar",   "title": "Remesas Familiares (mensual, MM US$)",
             "freq": "mensual", "var": "rem_fam", "anchor": "A42"},
            {"kind": "bar",   "title": "Llegada de No Residentes (mensual)",
             "freq": "mensual", "var": "se_cc_numtur", "anchor": "L42"},
        ],
    },
    {
        "sheet": "tab_tasas",
        "title": "TASAS DE INTERÉS",
        "charts": [
            {"kind": "line", "title": "Corredor de Tasas — TPM, Activa, Pasiva",
             "freq": "mensual", "vars": ["tpm", "tia_prom_pond", "tip_prom_pond"], "anchor": "A2"},
            {"kind": "line", "title": "TPM, Interbancaria y Fed Funds",
             "freq": "mensual", "vars": ["tpm", "tib_rd", "ti_fond_fed"], "anchor": "L2"},
            {"kind": "line", "title": "Inflación: General, Subyacente",
             "freq": "mensual", "vars": ["ipc_yoy", "ipc_sub"], "anchor": "A22"},
            {"kind": "bar",  "title": "TPM nivel actual (%)",
             "freq": "mensual", "var": "tpm", "anchor": "L22"},
        ],
    },
    {
        "sheet": "tab_tipo_cambio",
        "title": "TIPO DE CAMBIO",
        "charts": [
            {"kind": "line", "title": "TCN diario — Compra y Venta",
             "freq": "diaria", "vars": ["tcn_compra", "tcn_venta"], "anchor": "A2"},
            {"kind": "bar",  "title": "TCN Venta mensual",
             "freq": "mensual", "var": "tcn_venta_m", "anchor": "L2"},
            {"kind": "line", "title": "CPI USA + IPC RD (inputs del TCR)",
             "freq": "mensual", "vars": ["cpi_usa", "ipc_rd"], "anchor": "A22"},
        ],
    },
    {
        "sheet": "tab_depositos_prest",
        "title": "DEPÓSITOS Y PRÉSTAMOS",
        "charts": [
            {"kind": "bar",  "title": "Base Monetaria Restringida (MM RD$)",
             "freq": "mensual", "var": "bmr_rd", "anchor": "A2"},
            {"kind": "line", "title": "Préstamos por Sector (saldos)",
             "freq": "mensual", "vars": ["prest_agr", "prest_ind", "prest_con",
                                          "prest_com", "prest_viv", "prest_cons", "prest_tar"],
             "anchor": "L2"},
            {"kind": "bar",  "title": "Préstamos en ME (MM US$)",
             "freq": "mensual", "var": "prest_me_usd_osd", "anchor": "A22"},
            {"kind": "line", "title": "AEN OSD (Activos vs Pasivos externos)",
             "freq": "mensual", "vars": ["aen_osd_act", "aen_osd_pas"], "anchor": "L22"},
            {"kind": "bar",  "title": "Depósitos del Gobierno (OSD, MM RD$)",
             "freq": "mensual", "var": "dep_sp_osd", "anchor": "A42"},
        ],
    },
    {
        "sheet": "tab_dep_remunerados",
        "title": "DEPÓSITOS REMUNERADOS CP",
        "charts": [
            {"kind": "bar",  "title": "Saldo diario (MM RD$)",
             "freq": "diaria", "var": "deprem_cp", "anchor": "A2"},
            {"kind": "line", "title": "Saldo + TCN Venta (referencia)",
             "freq": "diaria", "vars": ["deprem_cp", "tcn_venta"], "anchor": "L2"},
        ],
    },
    {
        "sheet": "tab_sector_electrico",
        "title": "SECTOR ELÉCTRICO",
        "charts": [
            {"kind": "line", "title": "Generación por tipo de combustible (GWh)",
             "freq": "mensual", "vars": ["el_gen_total", "el_gen_carbon", "el_gen_gas_nat",
                                          "el_gen_hidro", "el_gen_eolica", "el_gen_solar"], "anchor": "A2"},
            {"kind": "line", "title": "Precios combustibles (US$/MMBTU)",
             "freq": "mensual", "vars": ["el_pcomb_fo2", "el_pcomb_fo6", "el_pcomb_gas", "el_pcomb_carbon"],
             "anchor": "L2"},
            {"kind": "line", "title": "Pérdidas por distribuidora (fracción)",
             "freq": "mensual", "vars": ["el_perd_tot", "el_perd_edn", "el_perd_eds", "el_perd_ede"],
             "anchor": "A22"},
            {"kind": "line", "title": "Cobranzas por distribuidora (fracción)",
             "freq": "mensual", "vars": ["el_cobr_tot", "el_cobr_edn", "el_cobr_eds", "el_cobr_ede"],
             "anchor": "L22"},
        ],
    },
]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    print("Cargando catálogo + series desde Supabase...")
    cat = fetch_catalog()
    series = fetch_all_series()

    # Filtrar catálogo a vars con data
    cat_with_data = [v for v in cat if v["id"] in series and series[v["id"]]]
    print(f"Variables con datos: {len(cat_with_data)}")

    wb = Workbook()
    wb.remove(wb.active)  # quita la hoja default

    # 1. README
    fecha_corte = pd.Timestamp.today().strftime("%Y-%m-%d %H:%M")
    n_filas = sum(len(v) for v in series.values())
    build_readme(wb, fecha_corte, len(cat_with_data), n_filas)

    # 2. Catalogo
    build_catalog(wb, cat_with_data)

    # 3. Hojas de datos por frecuencia
    by_freq: dict[str, list[dict]] = defaultdict(list)
    for v in cat_with_data:
        f = v["frecuencia"]
        if f in ("diaria", "mensual", "trimestral", "anual"):
            by_freq[f].append(v)
        elif f == "semanal":
            by_freq["diaria"].append(v)

    col_idx_per_freq: dict[str, dict[str, int]] = {}
    data_sheet_name: dict[str, str] = {
        "diaria": "data_diaria", "mensual": "data_mensual",
        "trimestral": "data_trimestral", "anual": "data_anual",
    }
    data_n_rows: dict[str, int] = {}
    for freq, vars_ in by_freq.items():
        vars_ = sorted(vars_, key=lambda v: v["id"])
        sn = data_sheet_name[freq]
        cols = build_data_sheet(wb, sn, vars_, series, freq)
        col_idx_per_freq[freq] = cols
        if cols:
            ws = wb[sn]
            # Filas de datos = max_row - 1 (excluir solo la fila 1 de cabecera)
            data_n_rows[freq] = ws.max_row - 1
        print(f"  data_{freq}: {len(vars_)} vars, {data_n_rows.get(freq, 0)} fechas")

    # 4. Hojas de gráficos por pestaña
    for spec in TAB_SPECS:
        ws = wb.create_sheet(spec["sheet"])
        ws.cell(row=1, column=1, value=spec["title"]).font = Font(size=14, bold=True)
        for chart in spec["charts"]:
            freq = chart["freq"]
            src = wb[data_sheet_name[freq]]
            cols = col_idx_per_freq.get(freq, {})
            n_rows = data_n_rows.get(freq, 0)
            if n_rows == 0:
                continue
            kind = chart["kind"]
            if kind == "bar":
                add_bar_chart(ws, src, title=chart["title"], var_id=chart["var"],
                              col_idx=cols, n_data_rows=n_rows, anchor=chart["anchor"],
                              freq=freq)
            elif kind == "line":
                add_line_chart(ws, src, title=chart["title"], var_ids=chart["vars"],
                               col_idx=cols, n_data_rows=n_rows, anchor=chart["anchor"],
                               freq=freq)
            elif kind == "combo":
                add_combo_chart(ws, src, title=chart["title"],
                                bar_var=chart["bar_var"], line_var=chart["line_var"],
                                col_idx=cols, n_data_rows=n_rows, anchor=chart["anchor"],
                                freq=freq)
        # Set wider column to fit charts
        for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K",
                            "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V"]:
            ws.column_dimensions[col_letter].width = 12

    # 5. Hoja de fórmulas
    build_formulas(wb)

    # Save
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    size_mb = OUTPUT.stat().st_size / 1024 / 1024
    print(f"\n✓ Guardado: {OUTPUT}")
    print(f"  Tamaño: {size_mb:.2f} MB")
    print(f"  Hojas: {wb.sheetnames}")


if __name__ == "__main__":
    main()
